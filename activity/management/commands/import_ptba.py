import logging
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from activity.models import (
    PTBA, PTBAStatus, Composante, SousComposante, Activite,
    SousActivite, FundingSource, SousActiviteFunding,
)

logger = logging.getLogger(__name__)

# Column indices in the 44-column PTBA Excel format (0-based)
COL_COMPOSANTE_CODE = 0
COL_COMPOSANTE_NAME = 1
COL_SOUS_COMPOSANTE_CODE = 2
COL_SOUS_COMPOSANTE_NAME = 3
COL_ACTIVITE_CODE = 4
COL_ACTIVITE_NAME = 5
COL_SOUS_ACTIVITE_CODE = 6
COL_SOUS_ACTIVITE_NAME = 7
COL_INDICATEUR = 8
COL_PROVINCE = 9
COL_CATEGORY_CODE = 10
COL_CATEGORY_NAME = 11
COL_STRUCTURE = 12
COL_PASSATION = 13
COL_UNITE = 14
COL_QTE_TOTAL = 15
COL_QTE_T1 = 16
COL_QTE_T2 = 17
COL_QTE_T3 = 18
COL_QTE_T4 = 19
COL_COUT_UNITAIRE = 20
COL_BUDGET_T1 = 21
COL_BUDGET_T2 = 22
COL_BUDGET_T3 = 23
COL_BUDGET_T4 = 24
COL_BUDGET_TOTAL = 25
COL_FUNDING_BM = 26
COL_FUNDING_UE = 27
COL_FUNDING_FIDA = 28
COL_FUNDING_BAD = 29
COL_FUNDING_ENABEL = 30
COL_FUNDING_ETC = 31

FUNDING_COLUMNS = [
    (COL_FUNDING_BM, 'BM'),
    (COL_FUNDING_UE, 'UE'),
    (COL_FUNDING_FIDA, 'FIDA'),
    (COL_FUNDING_BAD, 'BAD'),
    (COL_FUNDING_ENABEL, 'ENABEL'),
    (COL_FUNDING_ETC, 'GOV'),
]


def safe_decimal(value, default=Decimal('0')):
    """Convert a cell value to Decimal, returning default on failure."""
    if value is None:
        return default
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return default


def safe_str(value):
    """Convert a cell value to string, returning empty string for None."""
    if value is None:
        return ''
    return str(value).strip()


class Command(BaseCommand):
    help = (
        'Import a PTBA from a 44-column Excel file. '
        'Row 2 = headers, Row 3+ = data. '
        'Usage: python manage.py import_ptba <excel_file> --ptba-code PTBA-2024-2025'
    )

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Path to the PTBA Excel file')
        parser.add_argument(
            '--ptba-code', type=str, required=True,
            help='Code for the PTBA (e.g., PTBA-2024-2025). Creates if not exists.',
        )
        parser.add_argument(
            '--ptba-name', type=str, default=None,
            help='Name for the PTBA (used only when creating a new PTBA).',
        )
        parser.add_argument(
            '--fiscal-year-start', type=str, default=None,
            help='Fiscal year start date (YYYY-MM-DD). Used only when creating.',
        )
        parser.add_argument(
            '--fiscal-year-end', type=str, default=None,
            help='Fiscal year end date (YYYY-MM-DD). Used only when creating.',
        )
        parser.add_argument(
            '--sheet', type=str, default=None,
            help='Sheet name to read. Defaults to the active sheet.',
        )
        parser.add_argument(
            '--data-start-row', type=int, default=3,
            help='First data row (1-based). Default: 3 (row 2 = headers).',
        )
        parser.add_argument(
            '--dry-run', action='store_true',
            help='Parse and report without writing to the database.',
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl is required. Install with: pip install openpyxl")

        excel_file = options['excel_file']
        ptba_code = options['ptba_code']
        dry_run = options['dry_run']
        data_start_row = options['data_start_row']

        self.stdout.write(f"Loading Excel file: {excel_file}")

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
        except Exception as e:
            raise CommandError(f"Failed to open Excel file: {e}")

        if options['sheet']:
            if options['sheet'] not in wb.sheetnames:
                raise CommandError(
                    f"Sheet '{options['sheet']}' not found. "
                    f"Available: {wb.sheetnames}"
                )
            ws = wb[options['sheet']]
        else:
            ws = wb.active

        self.stdout.write(f"Using sheet: {ws.title}")
        self.stdout.write(f"Total rows: {ws.max_row}, Total columns: {ws.max_column}")

        # Pre-load funding sources
        funding_sources = {}
        for _, fs_code in FUNDING_COLUMNS:
            fs = FundingSource.objects.filter(code=fs_code).first()
            if fs:
                funding_sources[fs_code] = fs

        if not funding_sources:
            self.stdout.write(self.style.WARNING(
                "No funding sources found. Run 'seed_funding_sources' first."
            ))

        # Parse rows
        stats = {
            'composantes': 0,
            'sous_composantes': 0,
            'activites': 0,
            'sous_activites': 0,
            'funding_allocations': 0,
            'total_budget': Decimal('0'),
            'rows_processed': 0,
            'rows_skipped': 0,
        }

        current_composante = None
        current_sous_composante = None
        current_activite = None
        sort_order_composante = 0
        sort_order_sous_composante = 0
        sort_order_activite = 0
        sort_order_sous_activite = 0

        with transaction.atomic():
            # Get or create PTBA. A dry run performs the whole import and
            # rolls it back, so its summary is the one the real run prints.
            ptba = PTBA.objects.filter(code=ptba_code).first()
            if not ptba:
                ptba = PTBA.objects.create(
                    code=ptba_code,
                    name=options.get('ptba_name') or ptba_code,
                    fiscal_year_start=options.get('fiscal_year_start') or '2024-07-01',
                    fiscal_year_end=options.get('fiscal_year_end') or '2025-06-30',
                    status=PTBAStatus.DRAFT,
                )
                self.stdout.write(self.style.SUCCESS(f"Created PTBA: {ptba}"))
            else:
                self.stdout.write(f"Using existing PTBA: {ptba}")

            for row_num in range(data_start_row, ws.max_row + 1):
                row = [ws.cell(row=row_num, column=c + 1).value for c in range(max(32, ws.max_column))]

                # Skip completely empty rows
                if all(v is None for v in row[:8]):
                    stats['rows_skipped'] += 1
                    continue

                # Composante
                comp_code = safe_str(row[COL_COMPOSANTE_CODE])
                comp_name = safe_str(row[COL_COMPOSANTE_NAME])
                if comp_code and comp_name:
                    sort_order_composante += 1
                    current_composante, created = Composante.objects.update_or_create(
                        ptba=ptba, code=comp_code,
                        defaults={
                            'name': comp_name,
                            'sort_order': sort_order_composante,
                        },
                    )
                    if created:
                        stats['composantes'] += 1
                    sort_order_sous_composante = 0

                # SousComposante
                sc_code = safe_str(row[COL_SOUS_COMPOSANTE_CODE])
                sc_name = safe_str(row[COL_SOUS_COMPOSANTE_NAME])
                if sc_code and sc_name and current_composante:
                    sort_order_sous_composante += 1
                    current_sous_composante, created = SousComposante.objects.update_or_create(
                        composante=current_composante, code=sc_code,
                        defaults={
                            'name': sc_name,
                            'sort_order': sort_order_sous_composante,
                        },
                    )
                    if created:
                        stats['sous_composantes'] += 1
                    sort_order_activite = 0

                # Activite
                act_code = safe_str(row[COL_ACTIVITE_CODE])
                act_name = safe_str(row[COL_ACTIVITE_NAME])
                if act_name and current_sous_composante:
                    sort_order_activite += 1
                    activite_lookup = {'sous_composante': current_sous_composante}
                    if act_code:
                        activite_lookup['code'] = act_code
                    else:
                        activite_lookup['name'] = act_name
                    current_activite, created = Activite.objects.update_or_create(
                        **activite_lookup,
                        defaults={
                            'code': act_code,
                            'name': act_name,
                            'indicator_description': safe_str(row[COL_INDICATEUR]),
                            'province': safe_str(row[COL_PROVINCE]),
                            'implementing_structure': safe_str(row[COL_STRUCTURE]),
                            'procurement_method': safe_str(row[COL_PASSATION]),
                            'sort_order': sort_order_activite,
                        },
                    )
                    if created:
                        stats['activites'] += 1
                    sort_order_sous_activite = 0

                # SousActivite
                sa_code = safe_str(row[COL_SOUS_ACTIVITE_CODE])
                sa_name = safe_str(row[COL_SOUS_ACTIVITE_NAME])
                if sa_name and current_activite:
                    quantity_total = safe_decimal(row[COL_QTE_TOTAL])
                    quantity_t1 = safe_decimal(row[COL_QTE_T1])
                    quantity_t2 = safe_decimal(row[COL_QTE_T2])
                    quantity_t3 = safe_decimal(row[COL_QTE_T3])
                    quantity_t4 = safe_decimal(row[COL_QTE_T4])
                    unit_cost = safe_decimal(row[COL_COUT_UNITAIRE])
                    budget_t1 = safe_decimal(row[COL_BUDGET_T1])
                    budget_t2 = safe_decimal(row[COL_BUDGET_T2])
                    budget_t3 = safe_decimal(row[COL_BUDGET_T3])
                    budget_t4 = safe_decimal(row[COL_BUDGET_T4])
                    budget_total = safe_decimal(row[COL_BUDGET_TOTAL])

                    sort_order_sous_activite += 1
                    sa_lookup = {'activite': current_activite}
                    if sa_code:
                        sa_lookup['code'] = sa_code
                    else:
                        sa_lookup['name'] = sa_name

                    sa, created = SousActivite.objects.update_or_create(
                        **sa_lookup,
                        defaults={
                            'code': sa_code,
                            'name': sa_name,
                            'expense_category_code': safe_str(row[COL_CATEGORY_CODE]),
                            'expense_category': safe_str(row[COL_CATEGORY_NAME]),
                            'unit': safe_str(row[COL_UNITE]),
                            'quantity_total': quantity_total,
                            'quantity_t1': quantity_t1,
                            'quantity_t2': quantity_t2,
                            'quantity_t3': quantity_t3,
                            'quantity_t4': quantity_t4,
                            'unit_cost': unit_cost,
                            'budget_t1': budget_t1,
                            'budget_t2': budget_t2,
                            'budget_t3': budget_t3,
                            'budget_t4': budget_t4,
                            'budget_total': budget_total,
                            'sort_order': sort_order_sous_activite,
                        },
                    )
                    if created:
                        stats['sous_activites'] += 1
                    stats['total_budget'] += budget_total

                    # Funding allocations
                    for col_idx, fs_code in FUNDING_COLUMNS:
                        amount = safe_decimal(row[col_idx] if col_idx < len(row) else None)
                        if amount > 0 and fs_code in funding_sources:
                            SousActiviteFunding.objects.update_or_create(
                                sous_activite=sa,
                                funding_source=funding_sources[fs_code],
                                defaults={'amount': amount},
                            )
                            stats['funding_allocations'] += 1

                stats['rows_processed'] += 1

            if dry_run:
                transaction.set_rollback(True)

        # Report
        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS("Import Summary:"))
        self.stdout.write(f"  Rows processed: {stats['rows_processed']}")
        self.stdout.write(f"  Rows skipped (empty): {stats['rows_skipped']}")
        self.stdout.write(f"  Composantes created: {stats['composantes']}")
        self.stdout.write(f"  Sous-Composantes created: {stats['sous_composantes']}")
        self.stdout.write(f"  Activites created: {stats['activites']}")
        self.stdout.write(f"  Sous-Activites created: {stats['sous_activites']}")
        self.stdout.write(f"  Funding allocations: {stats['funding_allocations']}")
        self.stdout.write(f"  Total budget: {stats['total_budget']:,.2f} BIF")
        if dry_run:
            self.stdout.write(self.style.WARNING("  [DRY RUN] No changes were committed."))

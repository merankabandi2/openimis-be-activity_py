"""
UAT section F (docs/metier/activites-uat.md), rows F7 and F10: a line written
by import_ptba / import_ptba_revised edited afterwards through
updateSousActivite, and an execution reported on it.

The update input is built as openimis-fe-activity_js sousActiviteUpdatePayload
(src/utils/sous-activite.js) builds it: id, name and the changed text fields;
quarterly quantities, unit cost and the budgets recomputed from them only when
a quantity or the unit cost changed. The observed state is printed with the
prefix UATF.
"""
import datetime
import json
import os
import tempfile
import uuid
from decimal import Decimal
from io import StringIO

import openpyxl
from django.core.management import call_command

from core.models import MutationLog, Role, RoleRight
from core.models.openimis_graphql_test_case import openIMISGraphQLTestCase, BaseTestContext
from core.test_helpers import create_test_interactive_user

from activity.models import (
    PTBA, Composante, SousComposante, Activite, SousActivite, ActivityStatus,
)
from activity.services import QuarterlyExecutionService

ALL_ACTIVITY_RIGHTS = list(range(170005, 170018))


def _role_with_rights(name, rights):
    role = Role.objects.create(
        name=name, is_system=0, is_blocked=False, audit_user_id=-1,
        validity_from=datetime.datetime.now(),
    )
    for right_id in rights:
        RoleRight.objects.create(
            role=role, right_id=right_id, audit_user_id=-1,
            validity_from=datetime.datetime.now(),
        )
    return role


def _revised_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(3, 1, 'Composante')
    ws.cell(4, 1, 'Composante 1 : Filets sociaux UATF-BE')
    ws.cell(5, 2, 'Sous-composante 3.1 : Appui institutionnel UATF-BE')
    row = {5: 'UFB3.1.1', 6: 'Activite revisee UATF-BE', 7: 'Atelier de lancement UATF-BE', 8: 'atelier',
           9: 2, 10: 3, 11: 1000, 12: 1200, 13: 2000, 14: 3600, 15: 1500, 16: 2600, 17: 500, 18: 1000,
           19: 1600, 20: 'Révisé', 21: 'Hausse'}
    for col, value in row.items():
        ws.cell(6, col, value)
    wb.save(path)


def _fe_payload(sa, **changes):
    """updateSousActivite input for `sa` with `changes` (snake_case) applied."""
    payload = {'id': str(sa.id), 'name': changes.get('name', sa.name)}
    for field in ('code', 'unit', 'responsible'):
        if field in changes and changes[field] != getattr(sa, field):
            payload[field] = changes[field]
    quantity_fields = ('quantity_t1', 'quantity_t2', 'quantity_t3', 'quantity_t4', 'unit_cost')
    if any(f in changes and Decimal(str(changes[f])) != getattr(sa, f) for f in quantity_fields):
        values = {f: Decimal(str(changes.get(f, getattr(sa, f) or 0))) for f in quantity_fields}
        quarters = [values[f'quantity_t{i}'] for i in range(1, 5)]
        budgets = [q * values['unit_cost'] for q in quarters]
        payload.update(values)
        payload['quantity_total'] = sum(quarters)
        payload.update({f'budget_t{i}': budgets[i - 1] for i in range(1, 5)})
        payload['budget_total'] = sum(budgets)

    def camel(name):
        head, *rest = name.split('_')
        return head + ''.join(part.capitalize() for part in rest)
    return ' '.join(f'{camel(k)}: "{v}"' for k, v in payload.items())


def _state(sa):
    sa.refresh_from_db()
    return {k: str(getattr(sa, k)) for k in (
        'responsible', 'quantity_total', 'quantity_t1', 'unit_cost', 'budget_t1', 'budget_t2',
        'budget_t3', 'budget_t4', 'budget_total', 'budget_revised')}


class UatSectionFImportEditTests(openIMISGraphQLTestCase):

    @classmethod
    def setUpTestData(cls):
        role = _role_with_rights("UAT F PTBA Gestion", ALL_ACTIVITY_RIGHTS)
        cls.user = create_test_interactive_user(username="uat_f_gest_test", roles=[role.id])
        cls.token = BaseTestContext(user=cls.user).get_jwt()

    def _update(self, args):
        cmid = str(uuid.uuid4())
        response = self.query(
            f"""mutation {{ updateSousActivite(input: {{ {args} clientMutationId: "{cmid}" }}) {{ internalId }} }}""",
            headers={"HTTP_AUTHORIZATION": f"Bearer {self.token}"},
        )
        content = json.loads(response.content)
        log = MutationLog.objects.filter(client_mutation_id=cmid).first()
        return content, log

    def _f7_line(self):
        ptba = PTBA.objects.create(code='UATF-BE-F7', name='F7', fiscal_year_start='2024-07-01',
                                   fiscal_year_end='2025-06-30')
        comp = Composante.objects.create(ptba=ptba, code='UF1', name='C')
        sc = SousComposante.objects.create(composante=comp, code='UF1.1', name='SC')
        act = Activite.objects.create(sous_composante=sc, code='UF1.1.1', name='A')
        # Values written by import_ptba for the row with T1..T4 = 4000 and total 5000.
        sa = SousActivite.objects.create(
            activite=act, code='UF1.1.1.3', name='SA UATF trimestres incoherents', unit='unite',
            quantity_total=5, quantity_t1=1, quantity_t2=1, quantity_t3=1, quantity_t4=1, unit_cost=1000,
            budget_t1=1000, budget_t2=1000, budget_t3=1000, budget_t4=1000, budget_total=5000,
        )
        return act, sa

    def _f10_line(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'revise.xlsx')
            _revised_workbook(path)
            out = StringIO()
            call_command('import_ptba_revised', path, ptba_code='UATF-BE-F10', stdout=out)
        sa = SousActivite.objects.get(name='Atelier de lancement UATF-BE')
        return sa.activite, sa, out.getvalue()

    def test_f7_label_update_keeps_imported_budgets(self):
        act, sa = self._f7_line()
        content, log = self._update(_fe_payload(sa, responsible='Resp F7'))
        state = _state(sa)
        print('UATF F7 label', content.get('errors'), log and log.status, log and log.error, state)
        self.assertEqual(log.status, MutationLog.SUCCESS, log.error)
        self.assertEqual(state['responsible'], 'Resp F7')
        self.assertEqual([state[f'budget_t{i}'] for i in range(1, 5)], ['1000.00'] * 4)
        self.assertEqual(state['budget_total'], '5000.00')

    def test_f7_quantity_update_saves_recomputed_budgets(self):
        act, sa = self._f7_line()
        content, log = self._update(_fe_payload(sa, quantity_t1=2))
        state = _state(sa)
        print('UATF F7 quantity', content.get('errors'), log and log.status, log and log.error, state)
        self.assertEqual(log.status, MutationLog.SUCCESS, log.error)
        self.assertEqual(state['quantity_total'], '5.00')
        self.assertEqual([state[f'budget_t{i}'] for i in range(1, 5)],
                         ['2000.00', '1000.00', '1000.00', '1000.00'])
        self.assertEqual(state['budget_total'], '5000.00')

    def test_f7_budget_update_not_summing_to_total_is_refused(self):
        act, sa = self._f7_line()
        content, log = self._update(f'id: "{sa.id}" name: "{sa.name}" budgetTotal: "5000"')
        print('UATF F7 budget', content.get('errors'), log and log.status, log and log.error, _state(sa))
        self.assertEqual(log.status, MutationLog.ERROR)
        self.assertIn('do not sum to total', log.error)

    def test_f10_revised_line_quarters_are_filled(self):
        act, sa, out = self._f10_line()
        state = _state(sa)
        print('UATF F10 imported', state)
        self.assertEqual(state['budget_total'], '3600.00')
        self.assertEqual([state[f'budget_t{i}'] for i in range(1, 5)], ['900.00'] * 4)
        self.assertEqual(state['quantity_total'], '3.00')
        self.assertEqual(state['quantity_t1'], '0.75')

    def test_f10_label_update_saved(self):
        act, sa, out = self._f10_line()
        content, log = self._update(_fe_payload(sa, responsible='Resp F10'))
        state = _state(sa)
        print('UATF F10 label', content.get('errors'), log and log.status, log and log.error, state)
        self.assertEqual(log.status, MutationLog.SUCCESS, log.error)
        self.assertEqual(state['responsible'], 'Resp F10')
        self.assertEqual(state['budget_total'], '3600.00')

    def test_f10_quantity_update_saved(self):
        act, sa, out = self._f10_line()
        content, log = self._update(_fe_payload(sa, quantity_t1=1.75))
        state = _state(sa)
        print('UATF F10 quantity', content.get('errors'), log and log.status, log and log.error, state)
        self.assertEqual(log.status, MutationLog.SUCCESS, log.error)
        self.assertEqual(state['quantity_total'], '4.00')
        self.assertEqual([state[f'budget_t{i}'] for i in range(1, 5)],
                         ['2100.00', '900.00', '900.00', '900.00'])
        self.assertEqual(state['budget_total'], '4800.00')

    def test_f10_execution_rates_computed(self):
        act, sa, out = self._f10_line()
        act.status = ActivityStatus.EN_COURS
        act.save()
        execution = QuarterlyExecutionService.report(
            sa, 1, 2026, self.user, budget_engage=Decimal('1800'), budget_decaisse=Decimal('900'),
            resultats_realises=Decimal('2'), observations='F10',
        )
        execution.refresh_from_db()
        print('UATF F10 execution', {k: str(getattr(execution, k)) for k in (
            'budget_prevu', 'budget_engage', 'resultats_attendus', 'resultats_realises',
            'taux_engagement', 'taux_decaissement', 'taux_realisation')})
        self.assertEqual(execution.budget_prevu, Decimal('900'))
        self.assertEqual(execution.resultats_attendus, Decimal('0.75'))
        self.assertEqual(execution.taux_engagement, Decimal('200.00'))
        self.assertEqual(execution.taux_decaissement, Decimal('100.00'))
        self.assertEqual(execution.taux_realisation, Decimal('266.67'))

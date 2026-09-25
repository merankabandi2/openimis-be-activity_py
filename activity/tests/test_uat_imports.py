"""Import commands: dry run, quarterly split, duplicate lines and weekly-plan
matching (UAT activity 2026-09-25: ACT-S6, ACT-S14, DEF-F-03 to DEF-F-06)."""
import datetime
import io
import os
import tempfile
from decimal import Decimal

import openpyxl
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from activity.management.commands import import_ptba as ptba32
from activity.management.commands import import_ptba_revised as ptba21
from activity.management.commands import import_weekly_plan as weekly
from activity.models import (
    PTBA, Activite, Composante, SousActivite, SousComposante, WeeklyPlanEntry,
)
from activity.services import QuarterlyExecutionService
from activity.validation import validate_budget_consistency

QUARTERLY_BUDGETS = ('budget_t1', 'budget_t2', 'budget_t3', 'budget_t4')
QUARTERLY_QUANTITIES = ('quantity_t1', 'quantity_t2', 'quantity_t3', 'quantity_t4')


def _save(wb):
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def _row(ws, row_num, values):
    """values: {0-based column index: value}."""
    for col, value in values.items():
        ws.cell(row_num, col + 1, value)


def workbook_32(lines):
    """import_ptba format (data from row 3). lines: list of dicts
    {sa_code, sa_name, qty (T1..T4 list), unit_cost, budget (T1..T4 list), total}."""
    wb = openpyxl.Workbook()
    ws = wb.active
    _row(ws, 3, {ptba32.COL_COMPOSANTE_CODE: 'C1', ptba32.COL_COMPOSANTE_NAME: 'Composante un'})
    _row(ws, 4, {ptba32.COL_SOUS_COMPOSANTE_CODE: 'C1.1', ptba32.COL_SOUS_COMPOSANTE_NAME: 'SC un'})
    _row(ws, 5, {ptba32.COL_ACTIVITE_CODE: 'A1', ptba32.COL_ACTIVITE_NAME: 'Activite un'})
    r = 6
    for line in lines:
        q, b = line['qty'], line['budget']
        _row(ws, r, {
            ptba32.COL_SOUS_ACTIVITE_CODE: line['sa_code'],
            ptba32.COL_SOUS_ACTIVITE_NAME: line['sa_name'],
            ptba32.COL_QTE_TOTAL: sum(q),
            ptba32.COL_QTE_T1: q[0], ptba32.COL_QTE_T2: q[1],
            ptba32.COL_QTE_T3: q[2], ptba32.COL_QTE_T4: q[3],
            ptba32.COL_COUT_UNITAIRE: line['unit_cost'],
            ptba32.COL_BUDGET_T1: b[0], ptba32.COL_BUDGET_T2: b[1],
            ptba32.COL_BUDGET_T3: b[2], ptba32.COL_BUDGET_T4: b[3],
            ptba32.COL_BUDGET_TOTAL: line['total'],
        })
        r += 1
    return _save(wb)


def workbook_21(lines):
    """import_ptba_revised format (data from row 4). lines: list of dicts
    {act_code, sa_name, qty_init, qty_rev, uc, budget_init, budget_rev}."""
    wb = openpyxl.Workbook()
    ws = wb.active
    _row(ws, 4, {ptba21.COL_COMPOSANTE: 'Composante 1: Test'})
    r = 5
    for i, line in enumerate(lines):
        values = {
            ptba21.COL_ACT_CODE: line['act_code'],
            ptba21.COL_ACT_NAME: f"Activite {line['act_code']}",
            ptba21.COL_SA_NAME: line['sa_name'],
            ptba21.COL_UNITE: 'atelier',
            ptba21.COL_QTE_INITIALE: line['qty_init'],
            ptba21.COL_QTE_REVISEE: line.get('qty_rev'),
            ptba21.COL_COUT_UNITAIRE_INITIALE: line['uc'],
            ptba21.COL_COUT_UNITAIRE_REVISEE: line['uc'],
            ptba21.COL_BUDGET_INITIALE: line['budget_init'],
            ptba21.COL_BUDGET_REVISEE: line.get('budget_rev'),
        }
        if i == 0 or line.get('new_sc'):
            values[ptba21.COL_SC_CODE] = '1.1'
            values[ptba21.COL_SC_NAME] = 'Sous-composante 1.1'
        _row(ws, r, values)
        r += 1
    return _save(wb)


def workbook_weekly(rows, week_label='Etat de lieu au 23/03/2026'):
    """import_weekly_plan format (header row 1, data from row 3). rows: list of
    dicts {act_code, sa_name, responsible, etat, plan}."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, weekly.COL_WEEK_DATA_START + 1, week_label)
    ws.cell(1, weekly.COL_WEEK_DATA_START + 2, 'Planification du 23-27/03/2026')
    for i, row in enumerate(rows):
        _row(ws, 3 + i, {
            weekly.COL_ACT_CODE: row['act_code'],
            weekly.COL_ACT_NAME: f"Activite {row['act_code']}",
            weekly.COL_SA_NAME: row['sa_name'],
            weekly.COL_RESPONSABLE: row.get('responsible'),
            weekly.COL_WEEK_DATA_START: row.get('etat'),
            weekly.COL_WEEK_DATA_START + 1: row.get('plan'),
        })
    return _save(wb)


def run(command, *args, **options):
    out = io.StringIO()
    call_command(command, *args, stdout=out, **options)
    return out.getvalue()


def summary_value(output, label):
    for line in output.splitlines():
        if line.strip().startswith(label):
            return line.split(':', 1)[1].strip()
    raise AssertionError(f"{label!r} not in output:\n{output}")


def ptba_lines(code):
    return SousActivite.objects.filter(activite__sous_composante__composante__ptba__code=code)


class ImportDryRunTests(TestCase):
    """ACT-S14 / DEF-F-04: a dry run reports what the real run does and writes nothing."""

    LINES_32 = [
        {'sa_code': 'S1', 'sa_name': 'Ligne un', 'qty': [1, 1, 0, 0], 'unit_cost': 500,
         'budget': [500, 500, 0, 0], 'total': 1000},
        {'sa_code': 'S2', 'sa_name': 'Ligne deux', 'qty': [1, 2, 1, 1], 'unit_cost': 500,
         'budget': [500, 1000, 500, 500], 'total': 2500},
    ]
    LINES_21 = [
        {'act_code': 'A1', 'sa_name': 'Ligne un', 'qty_init': 2, 'qty_rev': 3, 'uc': 1000,
         'budget_init': 2000, 'budget_rev': 3000},
        {'act_code': 'A1', 'sa_name': 'Ligne deux', 'qty_init': 1, 'uc': 1000, 'budget_init': 1000},
    ]

    def _assert_same_summary(self, dry, real, labels):
        for label in labels:
            self.assertEqual(summary_value(dry, label), summary_value(real, label), label)

    def test_import_ptba_dry_run_on_new_code_counts_every_level_and_writes_nothing(self):
        path = workbook_32(self.LINES_32)
        dry = run('import_ptba', path, ptba_code='DRY-32-NEW', dry_run=True)
        self.assertFalse(PTBA.objects.filter(code='DRY-32-NEW').exists())
        self.assertEqual(summary_value(dry, 'Sous-Activites created'), '2')
        self.assertEqual(summary_value(dry, 'Total budget'), '3,500.00 BIF')
        real = run('import_ptba', path, ptba_code='DRY-32-NEW')
        self._assert_same_summary(dry, real, [
            'Composantes created', 'Sous-Composantes created', 'Activites created',
            'Sous-Activites created', 'Total budget'])
        self.assertEqual(ptba_lines('DRY-32-NEW').count(), 2)

    def test_import_ptba_dry_run_on_existing_ptba_reports_the_lines(self):
        path = workbook_32(self.LINES_32)
        run('import_ptba', path, ptba_code='DRY-32-OLD')
        before = ptba_lines('DRY-32-OLD').count()
        dry = run('import_ptba', path, ptba_code='DRY-32-OLD', dry_run=True)
        self.assertEqual(summary_value(dry, 'Total budget'), '3,500.00 BIF')
        self.assertEqual(ptba_lines('DRY-32-OLD').count(), before)

    def test_import_ptba_revised_dry_run_on_new_code_counts_every_level(self):
        path = workbook_21(self.LINES_21)
        dry = run('import_ptba_revised', path, ptba_code='DRY-21-NEW', dry_run=True)
        self.assertFalse(PTBA.objects.filter(code='DRY-21-NEW').exists())
        self.assertEqual(summary_value(dry, 'Sous-Activites created'), '2')
        self.assertEqual(summary_value(dry, 'Total budget (revised)'), '4,000.00 BIF')
        real = run('import_ptba_revised', path, ptba_code='DRY-21-NEW')
        self._assert_same_summary(dry, real, [
            'Composantes created', 'Sous-Composantes created', 'Activites created',
            'Sous-Activites created', 'Funding allocations', 'Total budget (initial)',
            'Total budget (revised)'])


class ImportRevisedQuarterlySplitTests(TestCase):
    """ACT-S6: the revised import fills T1-T4 so that they sum to the totals."""

    def test_new_line_is_split_evenly_and_is_consistent(self):
        path = workbook_21([{'act_code': 'A1', 'sa_name': 'Atelier', 'qty_init': 10, 'qty_rev': 12,
                             'uc': 1000, 'budget_init': 10000, 'budget_rev': 12000}])
        run('import_ptba_revised', path, ptba_code='SPLIT-NEW')
        sa = ptba_lines('SPLIT-NEW').get()
        self.assertEqual([getattr(sa, f) for f in QUARTERLY_BUDGETS], [Decimal('3000.00')] * 4)
        self.assertEqual([getattr(sa, f) for f in QUARTERLY_QUANTITIES], [Decimal('3.00')] * 4)
        self.assertIsNone(validate_budget_consistency(sa))

    def test_rounding_remainder_keeps_the_sum_exact(self):
        path = workbook_21([{'act_code': 'A1', 'sa_name': 'Atelier', 'qty_init': 1,
                             'uc': 1000, 'budget_init': 1000}])
        run('import_ptba_revised', path, ptba_code='SPLIT-ROUND')
        sa = ptba_lines('SPLIT-ROUND').get()
        self.assertEqual(sum(getattr(sa, f) for f in QUARTERLY_BUDGETS), Decimal('1000'))
        self.assertIsNone(validate_budget_consistency(sa))

    def test_reimport_keeps_the_shape_of_the_existing_split(self):
        path = workbook_21([{'act_code': 'A1', 'sa_name': 'Atelier', 'qty_init': 4,
                             'uc': 1000, 'budget_init': 4000}])
        run('import_ptba_revised', path, ptba_code='SPLIT-RE')
        sa = ptba_lines('SPLIT-RE').get()
        for f, v in zip(QUARTERLY_BUDGETS, (4000, 0, 0, 0)):
            setattr(sa, f, Decimal(v))
        for f, v in zip(QUARTERLY_QUANTITIES, (4, 0, 0, 0)):
            setattr(sa, f, Decimal(v))
        sa.save()
        path = workbook_21([{'act_code': 'A1', 'sa_name': 'Atelier', 'qty_init': 4, 'qty_rev': 6,
                             'uc': 1000, 'budget_init': 4000, 'budget_rev': 6000}])
        run('import_ptba_revised', path, ptba_code='SPLIT-RE')
        sa.refresh_from_db()
        self.assertEqual([getattr(sa, f) for f in QUARTERLY_BUDGETS],
                         [Decimal('6000.00'), 0, 0, 0])
        self.assertEqual(sa.quantity_t1, Decimal('6.00'))

    def test_execution_on_an_imported_line_gets_planned_budget_and_rates(self):
        path = workbook_21([{'act_code': 'A1', 'sa_name': 'Atelier', 'qty_init': 4,
                             'uc': 1000, 'budget_init': 4000}])
        run('import_ptba_revised', path, ptba_code='SPLIT-EXEC')
        sa = ptba_lines('SPLIT-EXEC').get()
        sa.activite.status = 'EN_COURS'
        sa.activite.save()
        from core.test_helpers import create_test_interactive_user
        user = create_test_interactive_user(username='split_exec_user')
        execution = QuarterlyExecutionService.report(
            sa, 1, 2026, user, budget_engage=Decimal('500'), budget_decaisse=Decimal('250'),
            resultats_realises=Decimal('1'),
        )
        self.assertEqual(execution.budget_prevu, Decimal('1000.00'))
        self.assertEqual(execution.taux_engagement, Decimal('50'))
        self.assertEqual(execution.taux_realisation, Decimal('100'))


class ImportRevisedDuplicateLineTests(TestCase):
    """DEF-F-06: a name repeated within one activity is refused, not merged."""

    def test_repeated_name_under_one_activity_aborts_the_import(self):
        path = workbook_21([
            {'act_code': 'A1', 'sa_name': 'Formation', 'qty_init': 1, 'uc': 5000, 'budget_init': 5000},
            {'act_code': 'A1', 'sa_name': 'Formation', 'qty_init': 1, 'uc': 7000, 'budget_init': 7000},
        ])
        with self.assertRaises(CommandError) as ctx:
            run('import_ptba_revised', path, ptba_code='DUP-21')
        self.assertIn("'Formation'", str(ctx.exception))
        self.assertIn('row 6 repeats row 5', str(ctx.exception))
        self.assertFalse(PTBA.objects.filter(code='DUP-21').exists())

    def test_same_name_under_two_activities_is_two_lines(self):
        path = workbook_21([
            {'act_code': 'A1', 'sa_name': 'Atelier', 'qty_init': 1, 'uc': 5000, 'budget_init': 5000},
            {'act_code': 'A2', 'sa_name': 'Atelier', 'qty_init': 1, 'uc': 7000, 'budget_init': 7000},
        ])
        out = run('import_ptba_revised', path, ptba_code='DUP-21-OK')
        self.assertEqual(ptba_lines('DUP-21-OK').count(), 2)
        self.assertEqual(summary_value(out, 'Total budget (initial)'), '12,000.00 BIF')


class ImportWeeklyPlanTests(TestCase):
    """DEF-F-03 / DEF-F-05: weekly rows match on activity code and name; the
    dry run reports what the real run creates."""

    @classmethod
    def setUpTestData(cls):
        cls.ptba = PTBA.objects.create(
            code='WEEKLY-PTBA', name='Weekly', fiscal_year_start=datetime.date(2026, 1, 1),
            fiscal_year_end=datetime.date(2026, 12, 31),
        )
        comp = Composante.objects.create(ptba=cls.ptba, code='1', name='C')
        sc = SousComposante.objects.create(composante=comp, code='1.1', name='SC')
        cls.act1 = Activite.objects.create(sous_composante=sc, code='V1.1', name='A1')
        cls.act2 = Activite.objects.create(sous_composante=sc, code='V2.2', name='A2')
        cls.sa1 = SousActivite.objects.create(activite=cls.act1, name='Atelier commun')
        cls.sa2 = SousActivite.objects.create(activite=cls.act2, name='Atelier commun')

    def test_same_name_under_two_activities_updates_each_line(self):
        path = workbook_weekly([
            {'act_code': 'V1.1', 'sa_name': 'Atelier commun', 'responsible': 'Resp V1.1',
             'etat': 'Etat V1.1', 'plan': 'Plan V1.1'},
            {'act_code': 'V2.2', 'sa_name': 'Atelier commun', 'responsible': 'Resp V2.2',
             'etat': 'Etat V2.2', 'plan': 'Plan V2.2'},
        ])
        out = run('import_weekly_plan', path, ptba_code='WEEKLY-PTBA')
        self.sa1.refresh_from_db()
        self.sa2.refresh_from_db()
        self.assertEqual(self.sa1.responsible, 'Resp V1.1')
        self.assertEqual(self.sa2.responsible, 'Resp V2.2')
        self.assertEqual(WeeklyPlanEntry.objects.get(sous_activite=self.sa1).status_description, 'Etat V1.1')
        self.assertEqual(WeeklyPlanEntry.objects.get(sous_activite=self.sa2).status_description, 'Etat V2.2')
        self.assertEqual(summary_value(out, 'Weekly entries created/updated'), '2')

    def test_dry_run_counts_new_weekly_lines_like_the_real_run(self):
        path = workbook_weekly([
            {'act_code': 'V1.1', 'sa_name': 'Suivi terrain', 'etat': 'E', 'plan': 'P'},
            {'act_code': 'V9.9', 'sa_name': 'Orpheline', 'etat': 'E', 'plan': 'P'},
        ])
        dry = run('import_weekly_plan', path, ptba_code='WEEKLY-PTBA', dry_run=True)
        self.assertFalse(SousActivite.objects.filter(name='Suivi terrain').exists())
        self.assertEqual(summary_value(dry, 'Sous-activites created (WEEKLY)'), '1')
        self.assertEqual(summary_value(dry, 'Sous-activites unmatched (no parent)'), '1')
        real = run('import_weekly_plan', path, ptba_code='WEEKLY-PTBA')
        for label in ('Sous-activites created (WEEKLY)', 'Sous-activites unmatched (no parent)',
                      'Weekly entries created/updated'):
            self.assertEqual(summary_value(dry, label), summary_value(real, label), label)
        self.assertTrue(SousActivite.objects.filter(name='Suivi terrain', activite=self.act1).exists())
